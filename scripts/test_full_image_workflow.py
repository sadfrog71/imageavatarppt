"""Zero-cost regression tests; solid fixtures are not ImageGen deliverables."""
import copy
import json
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation

from assemble_deck import build_deck
from audit_pptx import audit_pptx
from validate_plan import validate_plan


class FullImageWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.plan = json.loads((Path(__file__).resolve().parents[1] / 'manifest.example.json').read_text())
        self.plan['workflow']['style'].update(name='港华风格', confirmed=True, selection_evidence='TEST fixture only')
        base = self.plan['sections'][0]['slides'][0]
        pages = []
        for i in range(4):
            page = copy.deepcopy(base)
            page.update(id=f'p{i+1}', image=f'p{i+1}.png')
            pages.append(page)
        self.plan['sections'][0]['slides'] = pages
        self.plan['workflow']['preview']['slide_ids'] = ['p1', 'p2', 'p3']
        self.plan['output'] = 'test.pptx'

    def test_contract_rejects_bypass(self):
        self.assertEqual(validate_plan(self.plan), [])
        mutations = [
            lambda p: p['workflow']['style'].update(confirmed=False),
            lambda p: p['workflow']['preview'].update(slide_ids=['p1']),
            lambda p: p['workflow']['preview'].update(status='approved', approval_evidence=''),
            lambda p: p['typography'].update(prefer_editable_text=True),
            lambda p: p['image_generation'].update(local_text_overlay=True),
            lambda p: p.update(include_cover=True),
            lambda p: p['sections'][0]['slides'][0].update(title_render_mode='native'),
        ]
        for mutation in mutations:
            plan = copy.deepcopy(self.plan)
            mutation(plan)
            self.assertTrue(validate_plan(plan))
        self.plan['image_generation']['background_mode'] = 'soft_gradient'
        self.assertEqual(validate_plan(self.plan), [])

    def test_preview_gate_and_image_only_assembly(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / 'plan.json'
            path.write_text(json.dumps(self.plan))
            # Only preview images exist: preview must not read the fourth image.
            for page in self.plan['sections'][0]['slides'][:3]:
                Image.new('RGB', (2048, 1152), 'white').save(root / page['image'])
            preview, _ = build_deck(path, preview=True)
            self.assertEqual(len(Presentation(preview).slides), 3)
            with self.assertRaisesRegex(ValueError, 'approve'):
                build_deck(path)
            self.plan['workflow']['preview'].update(status='approved', approval_evidence='TEST approval only')
            path.write_text(json.dumps(self.plan))
            Image.new('RGB', (2048, 1152), 'white').save(root / 'p4.png')
            full, _ = build_deck(path)
            self.assertEqual(len(Presentation(full).slides), 4)
            report = audit_pptx(self.plan, full)
            self.assertTrue(report['passed'], report)
            prs = Presentation(full)
            prs.slides[0].shapes.add_textbox(0, 0, 1000, 1000).text = 'forbidden overlay'
            prs.save(full)
            self.assertFalse(audit_pptx(self.plan, full)['passed'])

    def test_short_deck_preview(self):
        self.plan['sections'][0]['slides'] = self.plan['sections'][0]['slides'][:2]
        self.plan['workflow']['preview']['slide_ids'] = ['p1', 'p2']
        self.assertEqual(validate_plan(self.plan), [])


if __name__ == '__main__':
    unittest.main()
